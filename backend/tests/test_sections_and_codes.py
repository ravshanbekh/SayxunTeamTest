"""
Algebra / geometriya bo'limlari (Rasch hisobotida) va takroriy test kodi (400, 500 emas).
"""
import pytest
from openpyxl import load_workbook

from app.services.rasch_report_service import (
    build_rasch_report, export_rasch_to_excel, export_rasch_to_pdf, item_labels_for, section_for_label,
)
from tests.test_rasch_report import _create_test, db  # noqa: F401 - fixture

GEOMETRY = ([str(i) for i in range(23, 31)] + ["33", "34", "35"] +
            [f"{i}{p}" for i in range(41, 45) for p in "ab"])


@pytest.mark.parametrize("label", item_labels_for("sertifikat"))
def test_section_mapping(label):
    expected = "Geometriya" if label in GEOMETRY else "Algebra"
    assert section_for_label(label) == expected


def test_section_sizes():
    labels = item_labels_for("sertifikat")
    geometry = [l for l in labels if section_for_label(l) == "Geometriya"]
    assert len(geometry) == 19          # 8 + 3 test, 4 yozma x 2 qism
    assert len(labels) - len(geometry) == 36
    for edge in ("23", "30", "33", "35", "41a", "44b"):
        assert section_for_label(edge) == "Geometriya"
    for edge in ("22", "31", "32", "36a", "40b", "45a", "45b", "1"):
        assert section_for_label(edge) == "Algebra"


async def test_report_sections_sum_to_total(db, tmp_path):  # noqa: F811
    test, _ = await _create_test(db, 60, seed=31)
    report = await build_rasch_report(db, test.id)
    assert report.section_totals() == {"Algebra": 36, "Geometriya": 19}
    labels = report.item_labels
    for row in report.rows:
        assert row.algebra + row.geometry == row.correct
        assert row.geometry == sum(v for v, l in zip(row.item_values, labels) if l in GEOMETRY)

    pdf = export_rasch_to_pdf(report, str(tmp_path / "r.pdf"))
    with open(pdf, "rb") as f:
        assert f.read(4) == b"%PDF"

    wb = load_workbook(export_rasch_to_excel(report, str(tmp_path / "r.xlsx")))
    ws = wb["Rasch natijalar"]
    header = [c.value for c in ws[1]]
    assert header[2:5] == ["To'g'ri", "Algebra", "Geometriya"]
    for r in range(2, ws.max_row + 1):
        assert ws.cell(r, 4).value + ws.cell(r, 5).value == ws.cell(r, 3).value
    assert [c.value for c in wb["Savollar tahlili"][1]][-1] == "Bo'lim"


async def test_prezident_report_has_no_sections(db):  # noqa: F811
    from app.models.test import Test
    test, _ = await _create_test(db, 5, seed=32)
    t = await db.get(Test, test.id)
    t.test_type = "prezident"
    await db.commit()
    report = await build_rasch_report(db, test.id)
    assert not report.has_sections
    assert "Algebra" not in report.headers()


async def test_duplicate_test_code_returns_400(db):  # noqa: F811
    from httpx import AsyncClient, ASGITransport
    from app.main import app
    from app.models.admin import AdminUser
    from app.utils.security import create_access_token, get_password_hash

    admin = AdminUser(username="admin_codes", password_hash=get_password_hash("secret1"), role="admin")
    db.add(admin)
    await db.commit()
    headers = {"Authorization": "Bearer " + create_access_token(
        {"sub": str(admin.id), "username": admin.username, "role": admin.role})}
    body = {"test_code": "mat1", "title": "Matematika", "test_type": "sertifikat",
            "answer_key": {"mcq_answers": {str(i): "A" for i in range(1, 36)}, "written_questions": {}}}

    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        r = await client.post("/api/v1/tests/", json=body, headers=headers)
        assert r.status_code == 201
        first_id = r.json()["id"]

        r = await client.post("/api/v1/tests/", json={**body, "test_code": "MAT1"}, headers=headers)
        assert r.status_code == 400
        assert "MAT1" in r.json()["detail"] and "mavjud" in r.json()["detail"]

        r = await client.post("/api/v1/tests/", json={**body, "test_code": "MAT2"}, headers=headers)
        assert r.status_code == 201
        second_id = r.json()["id"]

        # boshqa testning kodiga o'zgartirib bo'lmaydi
        r = await client.patch(f"/api/v1/tests/{second_id}", json={"test_code": "mat1"}, headers=headers)
        assert r.status_code == 400

        # o'z kodini qayta saqlash mumkin
        r = await client.patch(f"/api/v1/tests/{first_id}", json={"test_code": "MAT1", "title": "Yangi nom"},
                               headers=headers)
        assert r.status_code == 200 and r.json()["title"] == "Yangi nom"
