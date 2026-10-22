"""
Integratsiya: SQLite bazada test, talabalar va natijalar yaratib, Rasch
hisobotini (JSON / PDF / Excel) va bot uchun natijani tekshirish.
"""
import json
import os
import random
import uuid
from datetime import timedelta

import pytest
from openpyxl import load_workbook

from app.database import engine, Base, AsyncSessionLocal
from app.models.admin import AdminUser
from app.models.result import Result, MCQAnswer, WrittenAnswer
from app.models.session import TestSession
from app.models.test import Test, AnswerKey
from app.models.user import User
from app.services.rasch_report_service import (
    build_rasch_report, export_rasch_to_pdf, export_rasch_to_excel,
    get_cached_rasch_report, item_labels_for,
)
from app.utils.timer import now_uz

MCQ_KEY = {str(i): "ABCD"[i % 4] for i in range(1, 33)}
MCQ_KEY.update({"33": "E", "34": "F", "35": "A"})
WRITTEN_KEY = {str(i): {"a": str(i * 2), "b": str(i * 3)} for i in range(36, 46)}
NAMES = ["Humoyun", "Turayev", "Kamolova Shahrizoda", "Шахзода", "Ollohga oson🤲❤️", "No Name"]


async def _create_test(db, n_students, seed, end_time=None):
    rng = random.Random(seed)
    test = Test(
        test_code=f"MAT{seed}", title="Matematika sinov", test_type="sertifikat",
        is_active=True, end_time=end_time,
    )
    db.add(test)
    await db.flush()
    db.add(AnswerKey(test_id=test.id, mcq_answers=MCQ_KEY, written_questions=WRITTEN_KEY))

    difficulties = [rng.gauss(0.6, 1.2) for _ in range(55)]
    users = []
    for v in range(n_students):
        user = User(
            telegram_id=seed * 100000 + v, full_name=NAMES[v % len(NAMES)],
            surname=f"T{v}", region="Toshkent",
        )
        db.add(user)
        await db.flush()
        users.append(user)
        session = TestSession(
            user_id=user.id, test_id=test.id, session_token=uuid.uuid4().hex,
            expires_at=now_uz() + timedelta(hours=1), is_submitted=True,
        )
        db.add(session)
        await db.flush()

        theta = rng.gauss(0, 1)
        import math
        correct = [rng.random() < 1 / (1 + math.exp(-(theta - b))) for b in difficulties]
        result = Result(user_id=user.id, test_id=test.id, session_id=session.id,
                        mcq_score=sum(correct[:35]), written_score=sum(correct[35:]),
                        total_score=sum(correct), submitted_at=now_uz())
        db.add(result)
        await db.flush()
        for q in range(1, 36):
            right = MCQ_KEY[str(q)]
            answer = right if correct[q - 1] else ("B" if right != "B" else "C")
            db.add(MCQAnswer(result_id=result.id, question_number=q, student_answer=answer,
                             correct_answer=right, is_correct=correct[q - 1]))
        for k, q in enumerate(range(36, 46)):
            a_ok, b_ok = correct[35 + 2 * k], correct[36 + 2 * k]
            ans = {"a": WRITTEN_KEY[str(q)]["a"] if a_ok else "x",
                   "b": WRITTEN_KEY[str(q)]["b"] if b_ok else "y"}
            db.add(WrittenAnswer(result_id=result.id, question_number=q,
                                 student_answer=json.dumps(ans), score=int(a_ok) + int(b_ok),
                                 reviewed_at=now_uz()))
    await db.commit()
    return test, users


@pytest.fixture
async def db():
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.drop_all)
        await conn.run_sync(Base.metadata.create_all)
    async with AsyncSessionLocal() as session:
        yield session


async def test_report_rows_sorted_and_consistent(db):
    test, _ = await _create_test(db, 120, seed=1)
    report = await build_rasch_report(db, test.id)
    assert len(report.rows) == 120
    assert report.item_labels == item_labels_for("sertifikat")
    assert len(report.item_labels) == 55
    scores = [r.score for r in report.rows]
    assert scores == sorted(scores, reverse=True)
    assert [r.rank for r in report.rows] == list(range(1, 121))
    for r in report.rows:
        assert r.correct == sum(r.item_values)
        assert r.score >= 0
    # Bir xil to'g'ri javob -> bir xil Rasch ball (bitta variant)
    by_correct = {}
    for r in report.rows:
        by_correct.setdefault(r.correct, set()).add(r.score)
    assert all(len(v) == 1 for v in by_correct.values())


async def test_pdf_and_excel_export(db, tmp_path):
    test, _ = await _create_test(db, 80, seed=2)
    report = await build_rasch_report(db, test.id)

    pdf_path = export_rasch_to_pdf(report, str(tmp_path / "rasch.pdf"))
    assert os.path.getsize(pdf_path) > 2000
    with open(pdf_path, "rb") as f:
        assert f.read(4) == b"%PDF"

    xlsx_path = export_rasch_to_excel(report, str(tmp_path / "rasch.xlsx"))
    wb = load_workbook(xlsx_path)
    ws = wb["Rasch natijalar"]
    assert [c.value for c in ws[1]] == ["№", "Ism familiya", "To'g'ri", "Algebra", "Geometriya", "Rash", "Rash (algebra)", "Rash (geometriya)", "Foiz", "Daraja", "1-Fan", "2-Fan", "Manba"]
    assert ws.max_row == 81
    assert ws["F2"].value == report.rows[0].score
    assert ws["D2"].value + ws["E2"].value == ws["C2"].value
    assert round(ws["G2"].value + ws["H2"].value, 2) == ws["F2"].value
    assert len(list(wb["Savollar tahlili"].iter_rows())) == 56
    assert wb["Javoblar matritsasi"].max_column == 5 + 55


async def test_empty_test_report(db, tmp_path):
    test, _ = await _create_test(db, 0, seed=3)
    report = await build_rasch_report(db, test.id)
    assert report.rows == []
    export_rasch_to_pdf(report, str(tmp_path / "empty.pdf"))
    export_rasch_to_excel(report, str(tmp_path / "empty.xlsx"))


async def test_unknown_test_returns_none(db):
    assert await build_rasch_report(db, uuid.uuid4()) is None


async def test_cache_invalidated_on_new_result(db):
    test, users = await _create_test(db, 40, seed=4)
    first = await get_cached_rasch_report(db, test.id)
    again = await get_cached_rasch_report(db, test.id)
    assert first is again
    from sqlalchemy import select
    ak = (await db.execute(select(AnswerKey).where(AnswerKey.test_id == test.id))).scalars().first()
    ak.mcq_answers = {**MCQ_KEY, "1": "D"}
    await db.commit()
    changed = await get_cached_rasch_report(db, test.id)
    assert changed is not first


async def test_api_endpoints(db):
    from httpx import AsyncClient, ASGITransport
    from app.main import app
    from app.utils.security import create_access_token, get_password_hash

    test, users = await _create_test(db, 60, seed=5)
    admin = AdminUser(username="admin_rasch", password_hash=get_password_hash("secret"))
    db.add(admin)
    await db.commit()
    token = create_access_token({"sub": str(admin.id), "username": admin.username, "role": admin.role})
    headers = {"Authorization": f"Bearer {token}"}

    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        r = await client.get(f"/api/v1/admin/reports/{test.id}/rasch", headers=headers)
        assert r.status_code == 200, r.text
        data = r.json()
        assert data["participants"] == 60
        assert len(data["item_analysis"]) == 55

        r = await client.get(f"/api/v1/admin/reports/{test.id}/rasch/pdf", headers=headers)
        assert r.status_code == 200
        assert r.content[:4] == b"%PDF"

        r = await client.get(f"/api/v1/admin/reports/{test.id}/rasch/excel", headers=headers)
        assert r.status_code == 200

        r = await client.get(f"/api/v1/admin/reports/{test.id}/rasch")
        assert r.status_code in (401, 403)

        # Bot: test vaqti cheklanmagan -> Rasch natija darhol ko'rinadi
        r = await client.get(f"/api/v1/results/user/{users[0].id}/test-code/{test.test_code}")
        assert r.status_code == 200
        body = r.json()
        assert body["rasch"] is not None
        assert body["rasch"]["participants"] == 60
        assert body["rasch_pending"] is False


async def test_bot_rasch_hidden_until_test_ends(db):
    from httpx import AsyncClient, ASGITransport
    from app.main import app

    test, users = await _create_test(db, 30, seed=6, end_time=now_uz() + timedelta(hours=2))
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        r = await client.get(f"/api/v1/results/user/{users[0].id}/test-code/{test.test_code}")
        body = r.json()
        assert body["rasch"] is None
        assert body["rasch_pending"] is True
