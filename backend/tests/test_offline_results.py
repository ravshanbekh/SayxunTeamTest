"""
Offline natijalar: Excel shablon, fayl tekshiruvi, yuklash va Rasch hisobotiga
onlayn natijalar bilan birga qo'shilishi.
"""
import random
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.models.offline import OfflineResult
from app.models.test import Test
from app.services.offline_service import (
    build_template, parse_offline_excel, replace_offline_results, OfflineImportError,
)
from app.services.rasch_report_service import build_rasch_report, item_labels_for, get_cached_rasch_report
from app.services.rasch_service import compute_rasch
from tests.test_rasch_report import _create_test, db  # noqa: F401 - fixture

LABELS = item_labels_for("sertifikat")


def _workbook_bytes(rows, header=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Natijalar"
    ws.append(header or (["№", "Ism familiya", "Viloyat"] + LABELS))
    for row in rows:
        ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _random_answers(rng, p=0.5, n=55):
    return [1 if rng.random() < p else 0 for _ in range(n)]


# ---------------------------------------------------------------------------
# Shablon
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("test_type,count", [("sertifikat", 55), ("prezident", 40)])
def test_template_structure(test_type, count):
    test = Test(title="Matematika", test_code="MAT1", test_type=test_type)
    wb = load_workbook(BytesIO(build_template(test, test_type)))
    ws = wb["Natijalar"]
    header = [c.value for c in ws[1]]
    assert header[:3] == ["№", "Ism familiya", "Viloyat"]
    assert header[3:] == item_labels_for(test_type)
    assert len(header) == 3 + count
    assert "Yo'riqnoma" in wb.sheetnames
    assert ws.data_validations.dataValidation  # 0/1 tanlash ro'yxati


def test_empty_template_is_rejected():
    test = Test(title="Matematika", test_code="MAT1", test_type="sertifikat")
    with pytest.raises(OfflineImportError) as e:
        parse_offline_excel(build_template(test, "sertifikat"), "sertifikat")
    assert "topilmadi" in e.value.message


def test_filled_template_roundtrip():
    test = Test(title="Matematika", test_code="MAT1", test_type="sertifikat")
    wb = load_workbook(BytesIO(build_template(test, "sertifikat")))
    ws = wb["Natijalar"]
    rng = random.Random(1)
    expected = []
    for r in range(2, 12):
        answers = _random_answers(rng)
        expected.append(answers)
        ws.cell(row=r, column=2, value=f"Talaba {r}")
        ws.cell(row=r, column=3, value="Toshkent")
        for k, v in enumerate(answers):
            ws.cell(row=r, column=4 + k, value=v)
    buf = BytesIO()
    wb.save(buf)
    rows = parse_offline_excel(buf.getvalue(), "sertifikat")
    assert len(rows) == 10
    assert [r.answers for r in rows] == expected
    assert rows[0].full_name == "Talaba 2" and rows[0].region == "Toshkent"


# ---------------------------------------------------------------------------
# Fayl tekshiruvi
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("value,expected", [
    (1, 1), (0, 0), (1.0, 1), (0.0, 0), ("1", 1), ("0", 0), (" 1 ", 1), (None, 0), ("", 0), (True, 1),
])
def test_accepted_cell_values(value, expected):
    row = [1, "Ali Valiyev", ""] + [value] + [0] * 54
    rows = parse_offline_excel(_workbook_bytes([row]), "sertifikat")
    assert rows[0].answers[0] == expected


@pytest.mark.parametrize("value", [2, -1, 0.5, "A", "ha", "x", "11"])
def test_invalid_cell_values(value):
    row = [1, "Ali Valiyev", ""] + [1] * 10 + [value] + [0] * 44
    with pytest.raises(OfflineImportError) as e:
        parse_offline_excel(_workbook_bytes([row]), "sertifikat")
    assert e.value.errors and "11-savol" in e.value.errors[0] and "2-qator" in e.value.errors[0]


def test_missing_name_is_error():
    with pytest.raises(OfflineImportError) as e:
        parse_offline_excel(_workbook_bytes([[1, "", ""] + [1] * 55]), "sertifikat")
    assert "ism familiya" in e.value.errors[0]


def test_missing_columns():
    header = ["№", "Ism familiya"] + LABELS[:-2]
    with pytest.raises(OfflineImportError) as e:
        parse_offline_excel(_workbook_bytes([], header=header), "sertifikat")
    assert "45a" in e.value.message and "45b" in e.value.message


def test_missing_name_column():
    with pytest.raises(OfflineImportError):
        parse_offline_excel(_workbook_bytes([], header=["№", "X"] + LABELS), "sertifikat")


def test_blank_rows_skipped_and_columns_reordered():
    header = ["Viloyat"] + list(reversed(LABELS)) + ["Ism familiya"]
    answers = [1] * 20 + [0] * 35
    rows = [["Samarqand"] + list(reversed(answers)) + ["Bekzod"], [None] * 57, [None] * 57]
    parsed = parse_offline_excel(_workbook_bytes(rows, header=header), "sertifikat")
    assert len(parsed) == 1 and parsed[0].answers == answers and parsed[0].region == "Samarqand"


def test_q_prefixed_headers_from_regular_export():
    header = ["Talaba", "Correct"] + [f"Q{i}" for i in range(1, 36)] + [f"Q{i}{p}" for i in range(36, 46) for p in "ab"]
    parsed = parse_offline_excel(_workbook_bytes([["Ali", 3] + [1, 1, 1] + [0] * 52], header=header), "sertifikat")
    assert sum(parsed[0].answers) == 3


def test_prezident_40_items():
    header = ["№", "Ism familiya"] + [str(i) for i in range(1, 41)]
    parsed = parse_offline_excel(_workbook_bytes([[1, "Ali"] + [1] * 40], header=header), "prezident")
    assert parsed[0].answers == [1] * 40


def test_not_excel_file():
    with pytest.raises(OfflineImportError):
        parse_offline_excel(b"bu excel emas", "sertifikat")


def test_many_errors_are_capped():
    rows = [[i, f"T{i}", ""] + ["x"] * 55 for i in range(1, 30)]
    with pytest.raises(OfflineImportError) as e:
        parse_offline_excel(_workbook_bytes(rows), "sertifikat")
    assert len(e.value.errors) <= 55


# ---------------------------------------------------------------------------
# Baza va Rasch hisobiga qo'shilish
# ---------------------------------------------------------------------------
async def test_offline_rows_join_rasch_report(db):  # noqa: F811
    test, _ = await _create_test(db, 60, seed=21)
    online = await build_rasch_report(db, test.id)
    online_matrix = [r.item_values for r in online.rows]

    rng = random.Random(5)
    offline_rows = parse_offline_excel(_workbook_bytes(
        [[i, f"Offline {i}", "Buxoro"] + _random_answers(rng, p=0.4) for i in range(1, 41)]), "sertifikat")
    imported, replaced = await replace_offline_results(db, test.id, offline_rows)
    assert (imported, replaced) == (40, 0)

    report = await build_rasch_report(db, test.id)
    assert len(report.rows) == 100
    assert report.source_counts() == {"online": 60, "offline": 40}
    assert {r.name for r in report.rows if r.source == "offline"} == {f"Offline {i}" for i in range(1, 41)}
    assert all(r.user_id is None for r in report.rows if r.source == "offline")

    # Rasch bitta umumiy matritsada hisoblangan: mustaqil hisob bilan bir xil
    expected = compute_rasch(online_matrix + [r.answers for r in offline_rows])
    expected_by_raw = dict(zip(expected.raw_scores, expected.scores))
    for row in report.rows:
        assert row.score == expected_by_raw[row.correct]
    assert [r.rank for r in report.rows] == list(range(1, 101))


async def test_reupload_replaces_previous(db):  # noqa: F811
    test, _ = await _create_test(db, 10, seed=22)
    rng = random.Random(6)
    first = parse_offline_excel(_workbook_bytes(
        [[i, f"A{i}", ""] + _random_answers(rng) for i in range(1, 6)]), "sertifikat")
    second = parse_offline_excel(_workbook_bytes(
        [[i, f"B{i}", ""] + _random_answers(rng) for i in range(1, 4)]), "sertifikat")
    await replace_offline_results(db, test.id, first)
    imported, replaced = await replace_offline_results(db, test.id, second)
    assert (imported, replaced) == (3, 5)
    names = [o.full_name for o in (await db.execute(select(OfflineResult))).scalars().all()]
    assert sorted(names) == ["B1", "B2", "B3"]


async def test_cache_refreshes_after_offline_upload(db):  # noqa: F811
    test, _ = await _create_test(db, 20, seed=23)
    before = await get_cached_rasch_report(db, test.id)
    rows = parse_offline_excel(_workbook_bytes([[1, "Yangi", ""] + [1] * 55]), "sertifikat")
    await replace_offline_results(db, test.id, rows)
    after = await get_cached_rasch_report(db, test.id)
    assert after is not before and len(after.rows) == 21


async def test_offline_deleted_with_test(db):  # noqa: F811
    test, _ = await _create_test(db, 5, seed=24)
    rows = parse_offline_excel(_workbook_bytes([[1, "X", ""] + [1] * 55]), "sertifikat")
    await replace_offline_results(db, test.id, rows)
    await db.delete(await db.get(Test, test.id))
    await db.commit()
    assert (await db.execute(select(OfflineResult))).scalars().all() == []


async def test_offline_api_flow(db):  # noqa: F811
    from httpx import AsyncClient, ASGITransport
    from app.main import app
    from app.models.admin import AdminUser
    from app.utils.security import create_access_token, get_password_hash

    test, users = await _create_test(db, 30, seed=25)
    admin = AdminUser(username="admin_offline", password_hash=get_password_hash("secret1"), role="admin")
    db.add(admin)
    await db.commit()
    headers = {"Authorization": "Bearer " + create_access_token(
        {"sub": str(admin.id), "username": admin.username, "role": admin.role})}
    base = f"/api/v1/admin/tests/{test.id}/offline"

    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        r = await client.get(f"{base}/template", headers=headers)
        assert r.status_code == 200
        wb = load_workbook(BytesIO(r.content))
        ws = wb["Natijalar"]
        for k in range(5):
            ws.cell(row=2 + k, column=2, value=f"Qog'oz {k}")
            for c in range(55):
                ws.cell(row=2 + k, column=4 + c, value=1 if c < 10 + 5 * k else 0)
        buf = BytesIO()
        wb.save(buf)

        files = {"file": ("natija.xlsx", buf.getvalue(),
                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        r = await client.post(f"{base}/upload", headers=headers, files=files)
        assert r.status_code == 200, r.text
        assert r.json()["imported"] == 5

        r = await client.get(base, headers=headers)
        assert r.json()["count"] == 5
        assert [x["correct"] for x in r.json()["results"]] == [10, 15, 20, 25, 30]

        r = await client.get(f"/api/v1/admin/reports/{test.id}/rasch", headers=headers)
        data = r.json()
        assert data["participants"] == 35
        assert data["source_counts"] == {"online": 30, "offline": 5}

        # Bot: onlayn talaba o'rni endi 35 kishi ichida
        r = await client.get(f"/api/v1/results/user/{users[0].id}/test-code/{test.test_code}")
        assert r.json()["rasch"]["participants"] == 35

        bad = {"file": ("bad.xlsx", _workbook_bytes([[1, "Ali", ""] + ["A"] + [0] * 54]), "application/octet-stream")}
        r = await client.post(f"{base}/upload", headers=headers, files=bad)
        assert r.status_code == 400
        assert r.json()["detail"]["errors"]

        r = await client.delete(base, headers=headers)
        assert r.json()["deleted"] == 5

        r = await client.get(f"{base}/template")
        assert r.status_code in (401, 403)
