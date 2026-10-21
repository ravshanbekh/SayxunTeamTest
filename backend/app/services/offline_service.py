"""
Offline natijalar: Excel shablon yaratish, o'qituvchi to'ldirgan faylni
tekshirish va bazaga saqlash. Offline natijalar Rasch hisobida onlayn
natijalar bilan birga ishlatiladi.
"""
from dataclasses import dataclass
from io import BytesIO
from typing import List, Optional, Tuple
from uuid import UUID

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import select, delete, func
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.offline import OfflineResult
from app.services.rasch_report_service import item_labels_for

MAX_ROWS = 5000
MAX_ERRORS = 50
TEMPLATE_ROWS = 300

NAME_HEADERS = {"ism familiya", "ism-familiya", "ism", "f.i.sh", "fish", "fio", "talaba"}
REGION_HEADERS = {"viloyat", "hudud", "region", "maktab"}


@dataclass
class OfflineRow:
    full_name: str
    region: Optional[str]
    answers: List[int]


class OfflineImportError(Exception):
    def __init__(self, message: str, errors: Optional[List[str]] = None):
        super().__init__(message)
        self.message = message
        self.errors = errors or []


# ---------------------------------------------------------------------------
# Shablon
# ---------------------------------------------------------------------------
def build_template(test, test_type: str) -> bytes:
    labels = item_labels_for(test_type)
    wb = Workbook()
    ws = wb.active
    ws.title = "Natijalar"

    header_fill = PatternFill(start_color="ADD8E6", end_color="ADD8E6", fill_type="solid")
    ws.append(["№", "Ism familiya", "Viloyat"] + labels)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    for i in range(1, TEMPLATE_ROWS + 1):
        ws.cell(row=i + 1, column=1, value=i)

    first_item_col = get_column_letter(4)
    last_item_col = get_column_letter(3 + len(labels))
    dv = DataValidation(type="list", formula1='"0,1"', allow_blank=True,
                        showErrorMessage=True, errorTitle="Noto'g'ri qiymat",
                        error="Faqat 0 (xato) yoki 1 (to'g'ri) kiriting")
    dv.add(f"{first_item_col}2:{last_item_col}{TEMPLATE_ROWS + 1}")
    ws.add_data_validation(dv)

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 16
    for col in range(4, 4 + len(labels)):
        ws.column_dimensions[get_column_letter(col)].width = 5
    ws.freeze_panes = "D2"

    info = wb.create_sheet("Yo'riqnoma")
    lines = [
        f"Test: {test.title} (kod: {test.test_code})",
        "",
        "1. 'Natijalar' varag'ida har bir qatorga bitta talaba yoziladi.",
        "2. 'Ism familiya' majburiy, 'Viloyat' ixtiyoriy. '№' ustuni hisobga olinmaydi.",
        "3. Har bir savol ustuniga: 1 - to'g'ri javob, 0 - xato javob.",
        "   Bo'sh katak 0 (javob berilmagan) deb hisoblanadi.",
    ]
    if test_type != "prezident":
        lines.append("4. 1-35 - test savollari; 36a-45b - yozma savollarning a) va b) qismlari (har biri alohida 0/1).")
    else:
        lines.append("4. 1-40 - test savollari.")
    lines += [
        "5. Ustun nomlarini o'zgartirmang. Ortiqcha bo'sh qatorlar e'tiborga olinmaydi.",
        "6. Admin panelda 'Offline natija yuklash' tugmasi orqali faylni yuklang.",
        "   Qayta yuklansa, shu testning oldingi offline natijalari yangisi bilan almashtiriladi.",
        "7. Offline natijalar onlayn natijalar bilan birga Rasch hisobiga kiradi.",
        "",
        "Misol:  № | Ism familiya     | Viloyat  | 1 | 2 | 3 | ...",
        "        1 | Karimova Malika  | Toshkent | 1 | 0 | 1 | ...",
    ]
    for line in lines:
        info.append([line])
    info.column_dimensions["A"].width = 110

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Faylni o'qish va tekshirish
# ---------------------------------------------------------------------------
def _norm_header(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip().lower().replace(" ", "")
    if text.startswith("q") and len(text) > 1 and text[1].isdigit():
        text = text[1:]  # oddiy eksportdagi "Q1", "Q36a" ham qabul qilinadi
    return text


def _parse_cell(value) -> Optional[int]:
    """0/1 qiymat; bo'sh katak -> 0; noto'g'ri qiymat -> None."""
    if value is None:
        return 0
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, (int, float)):
        return int(value) if value in (0, 1) else None
    text = str(value).strip()
    if text == "":
        return 0
    if text in ("0", "1"):
        return int(text)
    if text in ("0.0", "1.0"):
        return int(float(text))
    return None


def parse_offline_excel(content: bytes, test_type: str) -> List[OfflineRow]:
    labels = item_labels_for(test_type)
    try:
        wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise OfflineImportError("Fayl o'qilmadi. Excel (.xlsx) formatidagi shablondan foydalaning.")

    ws = wb["Natijalar"] if "Natijalar" in wb.sheetnames else wb.worksheets[0]
    rows_iter = ws.iter_rows(values_only=True)
    header = next(rows_iter, None)
    if not header:
        raise OfflineImportError("Fayl bo'sh.")

    headers = [_norm_header(h) for h in header]
    name_col = next((i for i, h in enumerate(headers) if h in {x.replace(" ", "") for x in NAME_HEADERS}), None)
    region_col = next((i for i, h in enumerate(headers) if h in REGION_HEADERS), None)
    if name_col is None:
        raise OfflineImportError("'Ism familiya' ustuni topilmadi.")

    item_cols = {}
    for i, h in enumerate(headers):
        if h in labels and h not in item_cols:
            item_cols[h] = i
    missing = [label for label in labels if label not in item_cols]
    if missing:
        shown = ", ".join(missing[:15]) + (" ..." if len(missing) > 15 else "")
        raise OfflineImportError(f"Savol ustunlari topilmadi: {shown}. Shablonni qayta yuklab oling.")

    parsed: List[OfflineRow] = []
    errors: List[str] = []
    for row_num, row in enumerate(rows_iter, start=2):
        row = list(row) + [None] * (len(headers) - len(row))
        name = row[name_col]
        name = str(name).strip() if name is not None else ""
        values = [row[item_cols[label]] for label in labels]
        has_values = any(v not in (None, "") for v in values)
        if not name and not has_values:
            continue  # bo'sh qator
        if not name:
            errors.append(f"{row_num}-qator: ism familiya yozilmagan")
            continue
        answers = []
        for label, value in zip(labels, values):
            parsed_value = _parse_cell(value)
            if parsed_value is None:
                errors.append(f"{row_num}-qator ({name}), {label}-savol: '{value}' - faqat 0 yoki 1 bo'lishi kerak")
                parsed_value = 0
            answers.append(parsed_value)
        region = row[region_col] if region_col is not None else None
        region = str(region).strip() if region not in (None, "") else None
        parsed.append(OfflineRow(full_name=name[:255], region=region[:255] if region else None, answers=answers))
        if len(parsed) > MAX_ROWS:
            raise OfflineImportError(f"Juda ko'p qator (maksimal {MAX_ROWS}).")
        if len(errors) >= MAX_ERRORS:
            break

    if errors:
        raise OfflineImportError(f"Faylda {len(errors)} ta xato topildi. Tuzatib, qayta yuklang.", errors)
    if not parsed:
        raise OfflineImportError("Faylda birorta ham talaba natijasi topilmadi.")
    return parsed


# ---------------------------------------------------------------------------
# Baza
# ---------------------------------------------------------------------------
async def replace_offline_results(db: AsyncSession, test_id: UUID, rows: List[OfflineRow]) -> Tuple[int, int]:
    """Testning oldingi offline natijalarini yangisi bilan almashtirish. (yangi, o'chirilgan)."""
    old_count = await count_offline_results(db, test_id)
    await db.execute(delete(OfflineResult).where(OfflineResult.test_id == test_id))
    for row in rows:
        db.add(OfflineResult(test_id=test_id, full_name=row.full_name, region=row.region, answers=row.answers))
    await db.commit()
    return len(rows), old_count


async def count_offline_results(db: AsyncSession, test_id: UUID) -> int:
    stmt = select(func.count(OfflineResult.id)).where(OfflineResult.test_id == test_id)
    return (await db.execute(stmt)).scalar_one()


async def list_offline_results(db: AsyncSession, test_id: UUID) -> List[OfflineResult]:
    stmt = select(OfflineResult).where(OfflineResult.test_id == test_id).order_by(OfflineResult.created_at)
    return list((await db.execute(stmt)).scalars().all())


async def delete_offline_results(db: AsyncSession, test_id: UUID) -> int:
    count = await count_offline_results(db, test_id)
    await db.execute(delete(OfflineResult).where(OfflineResult.test_id == test_id))
    await db.commit()
    return count
