"""
Rasch natijalari hisoboti: milliy sertifikat ko'rinishidagi jadval
(№, Ism familiya, To'g'ri, Rash, Foiz, Daraja, 1-Fan, 2-Fan) - PDF va Excel.
"""
from dataclasses import dataclass
from typing import Dict, List, Optional
from uuid import UUID
from datetime import datetime
import json
import logging
import os
import time
import unicodedata

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (BaseDocTemplate, Frame, PageTemplate, NextPageTemplate,
                                Table, TableStyle, Paragraph, Spacer)
from sqlalchemy import select, func
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.offline import OfflineResult
from app.models.result import Result
from app.models.test import AnswerKey
from app.services.export_service import _load_export_data, _build_row_data
from app.config import settings
from app.services.rasch_service import compute_rasch, RaschResult, GRADE_THRESHOLDS

logger = logging.getLogger(__name__)


# Rasmdagi jadval ranglari
# Sayxun Team brend ranglari
BRAND_DARK = "#20242C"
BRAND_ORANGE = "#F58213"
BRAND_ORANGE_TINT = "#FFF1E3"
BRAND_MUTED = "#667085"
HEADER_COLOR = BRAND_DARK
GRID_COLOR = "#D0D5DD"
GRADE_COLORS = {
    "A+": "#90EE90",
    "A": "#ADFF2F",
    "B+": "#FFA500",
    "B": "#FFDAB9",
    "C+": "#FFE699",
    "C": "#FFFACD",
    None: "#F2F4F7",
}

HEADERS = ["№", "Ism familiya", "To'g'ri", "Rash", "Foiz", "Daraja", "1-Fan", "2-Fan"]

# Matematika sertifikat testi bo'limlari: geometriya savollari, qolganlari algebra
GEOMETRY_QUESTIONS = set(range(23, 31)) | {33, 34, 35} | set(range(41, 45))
SECTION_NAMES = ("Algebra", "Geometriya")


def section_for_label(label: str) -> str:
    """'27' yoki '42b' -> 'Geometriya', '5' yoki '36a' -> 'Algebra'."""
    digits = "".join(ch for ch in str(label) if ch.isdigit())
    return "Geometriya" if digits and int(digits) in GEOMETRY_QUESTIONS else "Algebra"


@dataclass
class RaschRow:
    rank: int
    result_id: Optional[UUID]
    user_id: Optional[UUID]
    name: str
    region: str
    correct: int
    item_values: List[int]
    theta: float
    score: float
    percent: Optional[float]
    grade: Optional[str]
    fan1: Optional[float]
    fan2: Optional[float]
    source: str = "online"  # "online" yoki "offline"
    algebra: Optional[int] = None     # algebra bo'limidagi to'g'ri javoblar
    geometry: Optional[int] = None    # geometriya bo'limidagi to'g'ri javoblar


@dataclass
class RaschReport:
    test: object
    test_type: str
    item_labels: List[str]
    rows: List[RaschRow]
    analysis: RaschResult

    @property
    def has_sections(self) -> bool:
        return self.test_type == "sertifikat"

    def section_totals(self) -> Dict[str, int]:
        totals = {name: 0 for name in SECTION_NAMES}
        for label in self.item_labels:
            totals[section_for_label(label)] += 1
        return totals

    def headers(self) -> List[str]:
        if not self.has_sections:
            return list(HEADERS)
        return HEADERS[:3] + list(SECTION_NAMES) + HEADERS[3:]

    def row_for_user(self, user_id) -> Optional[RaschRow]:
        for row in self.rows:
            if row.user_id == user_id:
                return row
        return None

    def source_counts(self) -> Dict[str, int]:
        online = sum(1 for row in self.rows if row.source == "online")
        return {"online": online, "offline": len(self.rows) - online}

    def grade_counts(self) -> Dict[str, int]:
        counts = {grade: 0 for _, grade in GRADE_THRESHOLDS}
        counts["-"] = 0
        for row in self.rows:
            counts[row.grade or "-"] += 1
        return counts


def item_labels_for(test_type: str) -> List[str]:
    """Topshiriq nomlari: sertifikat - 1..35 + 36a..45b, prezident - 1..40."""
    if test_type == "prezident":
        return [str(i) for i in range(1, 41)]
    labels = [str(i) for i in range(1, 36)]
    for i in range(36, 46):
        labels.extend([f"{i}a", f"{i}b"])
    return labels


async def build_rasch_report(db: AsyncSession, test_id: UUID) -> Optional[RaschReport]:
    """Test natijalaridan Rasch hisobotini tuzish."""
    t0 = time.time()
    test, results, users_dict, mcq_by_result, written_by_result, written_key = \
        await _load_export_data(db, test_id)
    if not test:
        return None

    test_type = getattr(test, "test_type", None) or "sertifikat"
    labels = item_labels_for(test_type)

    entries = []
    for result_record in results:
        user = users_dict.get(result_record.user_id)
        if not user:
            continue
        _, _, mcq_values, written_values = _build_row_data(
            user, result_record,
            mcq_by_result.get(result_record.id, []),
            written_by_result.get(result_record.id, []),
            written_key, test_type,
        )
        entries.append(dict(
            result_id=result_record.id, user_id=result_record.user_id, source="online",
            name=f"{user.full_name or ''} {user.surname or ''}".strip(),
            region=user.region or "", values=mcq_values + written_values,
        ))

    # Offline (qog'ozda) topshirilgan natijalar - umumiy Rasch hisobiga qo'shiladi
    stmt = select(OfflineResult).where(OfflineResult.test_id == test_id).order_by(OfflineResult.created_at)
    for offline in (await db.execute(stmt)).scalars().all():
        values = [1 if v else 0 for v in (offline.answers or [])][:len(labels)]
        values += [0] * (len(labels) - len(values))
        entries.append(dict(
            result_id=offline.id, user_id=None, source="offline",
            name=offline.full_name, region=offline.region or "", values=values,
        ))

    analysis = compute_rasch([e["values"] for e in entries], max_score=settings.RASCH_MAX_SCORE)

    rows = []
    for idx, entry in enumerate(entries):
        values = entry["values"]
        rows.append(RaschRow(
            rank=0,
            result_id=entry["result_id"],
            user_id=entry["user_id"],
            name=entry["name"],
            region=entry["region"],
            source=entry["source"],
            correct=analysis.raw_scores[idx],
            item_values=values,
            theta=analysis.thetas[idx],
            score=analysis.scores[idx],
            percent=analysis.percents[idx],
            grade=analysis.grades[idx],
            fan1=analysis.fan1[idx],
            fan2=analysis.fan2[idx],
        ))

    if test_type == "sertifikat":
        sections = [section_for_label(label) for label in labels]
        for row in rows:
            row.algebra = sum(v for v, sec in zip(row.item_values, sections) if sec == "Algebra")
            row.geometry = sum(v for v, sec in zip(row.item_values, sections) if sec == "Geometriya")

    rows.sort(key=lambda r: (-r.score, -r.correct, r.name.lower()))
    for rank, row in enumerate(rows, 1):
        row.rank = rank

    logger.info(f"[RASCH] {len(rows)} results analysed in {(time.time() - t0) * 1000:.0f}ms "
                f"(iterations={analysis.iterations}, converged={analysis.converged})")
    return RaschReport(test=test, test_type=test_type, item_labels=labels, rows=rows, analysis=analysis)


# ---------------------------------------------------------------------------
# Bot uchun kesh: natijalar yoki javoblar kaliti o'zgarmaguncha qayta hisoblamaslik
# ---------------------------------------------------------------------------
_report_cache: Dict[UUID, tuple] = {}


async def get_cached_rasch_report(db: AsyncSession, test_id: UUID) -> Optional[RaschReport]:
    stmt = select(func.count(Result.id), func.max(Result.submitted_at)).where(Result.test_id == test_id)
    count, last_submit = (await db.execute(stmt)).one()
    stmt = select(AnswerKey).where(AnswerKey.test_id == test_id)
    answer_key = (await db.execute(stmt)).scalars().first()
    key_sig = json.dumps(
        [answer_key.mcq_answers, answer_key.written_questions] if answer_key else None,
        sort_keys=True, default=str,
    )
    stmt = select(func.count(OfflineResult.id), func.max(OfflineResult.created_at)).where(
        OfflineResult.test_id == test_id)
    offline_count, offline_last = (await db.execute(stmt)).one()
    signature = (count, last_submit, offline_count, offline_last, key_sig)

    cached = _report_cache.get(test_id)
    if cached and cached[0] == signature:
        return cached[1]

    report = await build_rasch_report(db, test_id)
    if report is not None:
        _report_cache[test_id] = (signature, report)
    return report


# ---------------------------------------------------------------------------
# Formatlash
# ---------------------------------------------------------------------------
def fmt_num(value: Optional[float]) -> str:
    """74.75 -> '74.75', 75.0 -> '75.0', 73.7 -> '73.7' (rasmdagidek)."""
    if value is None:
        return "-"
    return str(round(float(value), 2))


def fmt_percent(value: Optional[float]) -> str:
    if value is None:
        return "-"
    if float(value).is_integer():
        return f"{int(value)}%"
    return f"{round(float(value), 2)}%"


def table_rows(report: RaschReport) -> List[List[str]]:
    rows = []
    for r in report.rows:
        sections = [str(r.algebra), str(r.geometry)] if report.has_sections else []
        rows.append([str(r.rank), r.name, str(r.correct)] + sections +
                    [fmt_num(r.score), fmt_percent(r.percent), r.grade or "-", fmt_num(r.fan1), fmt_num(r.fan2)])
    return rows


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
FONT_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "assets", "fonts")
_fonts_ready = None


def _register_fonts():
    """Kirill va lotin harflari uchun DejaVuSans; topilmasa Helvetica."""
    global _fonts_ready
    if _fonts_ready is None:
        try:
            pdfmetrics.registerFont(TTFont("DejaVuSans", os.path.join(FONT_DIR, "DejaVuSans.ttf")))
            pdfmetrics.registerFont(TTFont("DejaVuSans-Bold", os.path.join(FONT_DIR, "DejaVuSans-Bold.ttf")))
            _fonts_ready = ("DejaVuSans", "DejaVuSans-Bold")
        except Exception as e:
            logger.warning(f"[RASCH] DejaVuSans not available, using Helvetica: {e}")
            _fonts_ready = ("Helvetica", "Helvetica-Bold")
    return _fonts_ready


def _pdf_safe(text: str) -> str:
    """Emoji va shriftda yo'q belgilarni olib tashlash (PDFda kvadrat chiqmasligi uchun)."""
    cleaned = []
    for ch in text:
        if ord(ch) > 0xFFFF or unicodedata.category(ch) in ("So", "Cs", "Co", "Cn"):
            continue
        if 0xFE00 <= ord(ch) <= 0xFE0F or ch == "‍":
            continue
        cleaned.append(ch)
    result = "".join(cleaned).strip()
    return result or "-"


BRAND_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "assets", "brand")


def _brand_image(name: str) -> Optional[str]:
    path = os.path.join(BRAND_DIR, name)
    return path if os.path.exists(path) else None


def export_rasch_to_pdf(report: RaschReport, filepath: str) -> str:
    regular, bold = _register_fonts()
    os.makedirs(os.path.dirname(filepath) or ".", exist_ok=True)

    page_w, page_h = A4
    margin = 0.55 * inch
    first_band = 92      # birinchi sahifadagi logotipli sarlavha
    later_band = 40      # keyingi sahifalardagi ingichka sarlavha
    footer_h = 30
    logo_path = _brand_image("logo.png")
    watermark_path = _brand_image("watermark.png")
    title_text = _pdf_safe(f"{report.test.title}")
    generated = datetime.now().strftime("%Y-%m-%d %H:%M")

    def draw_footer(canvas, doc):
        canvas.setStrokeColor(colors.HexColor(BRAND_ORANGE))
        canvas.setLineWidth(1.2)
        canvas.line(margin, footer_h, page_w - margin, footer_h)
        canvas.setFont(regular, 7.5)
        canvas.setFillColor(colors.HexColor(BRAND_MUTED))
        canvas.drawString(margin, footer_h - 12, "Sayxun Team o'quv markazi · Rasch natijalari")
        canvas.drawRightString(page_w - margin, footer_h - 12, f"{doc.page}-bet")

    def draw_watermark(canvas):
        # Kontent ustidan juda xira chiziladi: rangli qatorlar ostida ham ko'rinadi
        if not watermark_path:
            return
        w = page_w * 0.78
        h = w * 0.29
        canvas.drawImage(watermark_path, (page_w - w) / 2, (page_h - h) / 2 - 20, width=w, height=h,
                         mask="auto", preserveAspectRatio=True)

    def first_page(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(colors.HexColor(BRAND_DARK))
        canvas.rect(0, page_h - first_band, page_w, first_band, stroke=0, fill=1)
        canvas.setFillColor(colors.HexColor(BRAND_ORANGE))
        canvas.rect(0, page_h - first_band - 3, page_w, 3, stroke=0, fill=1)
        if logo_path:
            canvas.drawImage(logo_path, margin, page_h - first_band + 14, width=136, height=64,
                             mask="auto", preserveAspectRatio=True, anchor="w")
        canvas.setFillColor(colors.white)
        canvas.setFont(bold, 15)
        canvas.drawRightString(page_w - margin, page_h - 40, "Rasch natijalari")
        canvas.setFont(regular, 9.5)
        canvas.setFillColor(colors.HexColor("#F7B26A"))
        canvas.drawRightString(page_w - margin, page_h - 56, title_text[:70])
        canvas.setFillColor(colors.HexColor("#C9CED6"))
        canvas.setFont(regular, 8)
        canvas.drawRightString(page_w - margin, page_h - 70, "Milliy sertifikat uslubida baholash")
        draw_footer(canvas, doc)
        canvas.restoreState()

    def later_page(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(colors.HexColor(BRAND_DARK))
        canvas.rect(0, page_h - later_band, page_w, later_band, stroke=0, fill=1)
        canvas.setFillColor(colors.HexColor(BRAND_ORANGE))
        canvas.rect(0, page_h - later_band - 2, page_w, 2, stroke=0, fill=1)
        if logo_path:
            canvas.drawImage(logo_path, margin, page_h - later_band + 6, width=62, height=28,
                             mask="auto", preserveAspectRatio=True, anchor="w")
        canvas.setFillColor(colors.white)
        canvas.setFont(regular, 8.5)
        canvas.drawRightString(page_w - margin, page_h - 24, f"{title_text[:60]} — Rasch natijalari")
        draw_footer(canvas, doc)
        canvas.restoreState()

    def page_end(canvas, doc):
        canvas.saveState()
        draw_watermark(canvas)
        canvas.restoreState()

    doc = BaseDocTemplate(
        filepath, pagesize=A4,
        title=f"{report.test.title} - Rasch natijalari",
        author="Sayxun Team o'quv markazi",
    )
    frame_w = page_w - 2 * margin
    first_frame = Frame(margin, footer_h + 10, frame_w, page_h - first_band - footer_h - 26, id="first",
                        leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
    later_frame = Frame(margin, footer_h + 10, frame_w, page_h - later_band - footer_h - 24, id="later",
                        leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
    doc.addPageTemplates([
        PageTemplate(id="First", frames=[first_frame], onPage=first_page, onPageEnd=page_end),
        PageTemplate(id="Later", frames=[later_frame], onPage=later_page, onPageEnd=page_end),
    ])

    styles = getSampleStyleSheet()
    note_style = ParagraphStyle("RaschNote", parent=styles["Normal"], fontName=regular,
                                fontSize=7, textColor=colors.HexColor(BRAND_MUTED), leading=9.5)
    section_style = ParagraphStyle("RaschSection", parent=styles["Normal"], fontName=bold,
                                   fontSize=9.5, textColor=colors.HexColor(BRAND_DARK), spaceAfter=6)

    analysis = report.analysis
    sources = report.source_counts()

    # Asosiy ma'lumotlar: yorliq + qiymat
    facts = [
        ("Test kodi", _pdf_safe(str(report.test.test_code))),
        ("Ishtirokchilar", f"{len(report.rows)}  (onlayn {sources['online']}, offline {sources['offline']})"),
        ("Topshiriqlar", str(len(report.item_labels)) + (
            f"  (algebra {report.section_totals()['Algebra']}, geometriya {report.section_totals()['Geometriya']})"
            if report.has_sections else "")),
        ("Sana", generated),
    ]
    facts_table = Table([[f[0] for f in facts], [f[1] for f in facts]],
                        colWidths=[62, 168, 176, frame_w - 406])
    facts_table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, 0), regular),
        ("FONTNAME", (0, 1), (-1, 1), bold),
        ("FONTSIZE", (0, 0), (-1, 0), 7),
        ("FONTSIZE", (0, 1), (-1, 1), 9),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor(BRAND_MUTED)),
        ("TEXTCOLOR", (0, 1), (-1, 1), colors.HexColor(BRAND_DARK)),
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(BRAND_ORANGE_TINT)),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, 0), 6),
        ("BOTTOMPADDING", (0, 1), (-1, 1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 0),
    ]))

    elements = [NextPageTemplate("Later"), facts_table, Spacer(1, 12)]

    if report.has_sections:
        totals = report.section_totals()
        header = ["№", "Ism familiya", f"To'g'ri\n({len(report.item_labels)})",
                  f"Algebra\n({totals['Algebra']})", f"Geometriya\n({totals['Geometriya']})",
                  "Rash", "Foiz", "Daraja", "1-Fan", "2-Fan"]
        col_widths = [22, 128, 38, 48, 58, 46, 50, 40, 43, 43]
    else:
        header = list(HEADERS)
        col_widths = [28, 165, 46, 52, 54, 48, 60, 60]
    data = [header] + [
        [row[0], _pdf_safe(row[1])] + row[2:] for row in table_rows(report)
    ]
    table = Table(data, colWidths=col_widths, repeatRows=1)

    style = [
        ("FONTNAME", (0, 0), (-1, 0), bold),
        ("FONTNAME", (0, 1), (-1, -1), regular),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(BRAND_DARK)),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("LINEBELOW", (0, 0), (-1, 0), 1.5, colors.HexColor(BRAND_ORANGE)),
        ("INNERGRID", (0, 1), (-1, -1), 0.6, colors.white),
        ("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor(GRID_COLOR)),
        ("ALIGN", (0, 0), (-1, -1), "LEFT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 2.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
        ("TOPPADDING", (0, 0), (-1, 0), 5),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 5),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
    ]
    for i, row in enumerate(report.rows, 1):
        style.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor(GRADE_COLORS[row.grade])))
    table.setStyle(TableStyle(style))
    elements.append(table)

    # Darajalar bo'yicha qisqa statistika
    counts = report.grade_counts()
    elements.append(Spacer(1, 16))
    elements.append(Paragraph("Darajalar bo'yicha taqsimot", section_style))
    summary = [list(counts.keys()), [str(v) for v in counts.values()]]
    summary_table = Table(summary, colWidths=[46] * len(counts))
    summary_style = [
        ("FONTNAME", (0, 0), (-1, 0), bold),
        ("FONTNAME", (0, 1), (-1, 1), regular),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("INNERGRID", (0, 0), (-1, -1), 0.6, colors.white),
        ("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor(GRID_COLOR)),
        ("BACKGROUND", (0, 1), (-1, 1), colors.white),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]
    for col, grade in enumerate(counts.keys()):
        color = GRADE_COLORS.get(None if grade == "-" else grade)
        summary_style.append(("BACKGROUND", (col, 0), (col, 0), colors.HexColor(color)))
    summary_table.setStyle(TableStyle(summary_style))
    summary_table.hAlign = "LEFT"
    elements.append(summary_table)

    rel = f"{analysis.reliability:.2f}" if analysis.reliability is not None else "-"
    elements.append(Spacer(1, 10))
    elements.append(Paragraph(
        "Ballar Rasch modeli (CML + WLE) asosida hisoblangan: T = 50 + 10·Z (shartli maksimal 75, undan yuqori ham bo'lishi mumkin). "
        "Darajalar: A+ ≥70, A ≥65, B+ ≥60, B ≥55, C+ ≥50, C ≥46. "
        "Foiz: A+/A – 100%, qolganlari ball/65. 1-Fan = foiz·93 + 11, 2-Fan = foiz·63 + 11 "
        f"(majburiy fan 11 ball). Ishonchlilik (reliability): {rel}.",
        note_style,
    ))
    for warning in analysis.warnings:
        elements.append(Paragraph(f"⚠ {warning}", note_style))

    doc.build(elements)
    return filepath


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
def export_rasch_to_excel(report: RaschReport, filepath: str) -> str:
    os.makedirs(os.path.dirname(filepath) or ".", exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Rasch natijalar"

    side = Side(style="thin", color=GRID_COLOR.lstrip("#"))
    border = Border(left=side, right=side, top=side, bottom=side)
    header_fill = PatternFill(start_color=HEADER_COLOR.lstrip("#"), end_color=HEADER_COLOR.lstrip("#"),
                              fill_type="solid")
    fills = {g: PatternFill(start_color=c.lstrip("#"), end_color=c.lstrip("#"), fill_type="solid")
             for g, c in GRADE_COLORS.items()}

    headers = report.headers()
    ws.append(headers + ["Manba"])
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True, color="FFFFFF")
        cell.border = border
        cell.alignment = Alignment(horizontal="left", vertical="center")

    for r in report.rows:
        sections = [r.algebra, r.geometry] if report.has_sections else []
        ws.append([r.rank, r.name, r.correct] + sections + [
            r.score,
            (r.percent / 100.0) if r.percent is not None else "-",
            r.grade or "-",
            r.fan1 if r.fan1 is not None else "-",
            r.fan2 if r.fan2 is not None else "-",
            "Offline" if r.source == "offline" else "Onlayn",
        ])
        row_idx = ws.max_row
        for cell in ws[row_idx]:
            cell.fill = fills[r.grade]
            cell.border = border
            cell.alignment = Alignment(horizontal="left", vertical="center")
        col = {name: i + 1 for i, name in enumerate(headers)}
        ws.cell(row=row_idx, column=col["Rash"]).number_format = "0.00"
        ws.cell(row=row_idx, column=col["Foiz"]).number_format = "0.00%"
        ws.cell(row=row_idx, column=col["1-Fan"]).number_format = "0.00"
        ws.cell(row=row_idx, column=col["2-Fan"]).number_format = "0.00"

    widths = {"№": 6, "Ism familiya": 34, "Algebra": 10, "Geometriya": 12, "Manba": 10}
    for i, name in enumerate(headers + ["Manba"], 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(name, 9)
    ws.freeze_panes = "A2"

    # Savollar tahlili
    ws2 = wb.create_sheet("Savollar tahlili")
    ws2.append(["Savol", "To'g'ri javoblar", "Foiz", "Qiyinlik (logit)", "Izoh"] +
               (["Bo'lim"] if report.has_sections else []))
    for cell in ws2[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True, color="FFFFFF")
        cell.border = border
    n = len(report.rows) or 1
    for label, item in zip(report.item_labels, report.analysis.items):
        ws2.append([
            label, item.correct_count, item.correct_count / n, round(item.difficulty, 3),
            "Chegaraviy (0.3 tuzatma)" if item.extreme else "",
        ] + ([section_for_label(label)] if report.has_sections else []))
        ws2.cell(row=ws2.max_row, column=3).number_format = "0.0%"
    for col, width in zip("ABCDEF", [8, 16, 10, 16, 26, 12]):
        ws2.column_dimensions[col].width = width

    # Talabalar bo'yicha 0/1 matritsa (tekshirish uchun)
    ws3 = wb.create_sheet("Javoblar matritsasi")
    ws3.append(["№", "Ism familiya", "To'g'ri", "θ (logit)", "Rash"] + report.item_labels)
    for cell in ws3[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True, color="FFFFFF")
    for r in report.rows:
        ws3.append([r.rank, r.name, r.correct, round(r.theta, 4), r.score] + list(r.item_values))
    ws3.column_dimensions["B"].width = 30

    wb.save(filepath)
    return filepath
